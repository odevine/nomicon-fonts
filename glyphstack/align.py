"""Aligns segmented glyph boxes to the known printed text so every box gets its character"""
import numpy as np
import openpyxl

SKIP_BOX = 7.0
SKIP_CHAR = 7.0
MERGE_PENALTY = 1.5


def card_text(xlsx, name):
    """Returns the rules text of a card from the Card text tab, without the Rules: and Flavor: labels"""
    ws = openpyxl.load_workbook(xlsx)["Card text"]
    for row in ws.iter_rows(values_only=True):
        if row[1] == name:
            lines = [l for l in row[4].split("\n") if l not in ("Rules:", "Flavor:")]
            return "\n".join(lines)
    raise KeyError(name)


def line_metrics(glyphs):
    """Adds baseline-relative top, bottom and width in x-height units to each glyph dict"""
    for line in {g["line"] for g in glyphs}:
        row = [g for g in glyphs if g["line"] == line]
        bottoms = np.array([g["y"] + g["h"] for g in row])
        base = float(np.median(bottoms))
        on_base = [g for g in row if abs(g["y"] + g["h"] - base) < 0.06 * np.median([g["h"] for g in row])]
        xh = float(np.median([g["h"] for g in on_base]))
        for g in row:
            g["top"] = (base - g["y"]) / xh
            g["bottom"] = (g["y"] + g["h"] - base) / xh
            g["width"] = g["w"] / xh
            g["xh"] = xh


# prototype (top, bottom) in x-height units for the shape classes of the text face
def _prior(c):
    if c in "aceimnorsuvwxz":
        return (1.0, 0.0)
    if c in "bdfhklt":
        return (1.45, 0.0) if c != "t" else (1.25, 0.0)
    if c in "gpqy":
        return (1.0, 0.45)
    if c == "j":
        return (1.45, 0.45)
    if c in ",;":
        return (0.2, 0.3)
    if c in ".:":
        return (0.2, 0.0)
    if c in "“”‘’\"'":
        return (1.45, -0.6)
    if c == "/":
        return (1.45, 0.35)
    if c.isupper() or c.isdigit():
        return (1.4, 0.0)
    return (1.0, 0.0)


class Protos:
    """Per-character expected (top, bottom, width), starting from class priors and refined from alignments"""

    def __init__(self, chars):
        self.p = {c: np.array([*_prior(c), np.nan]) for c in set(chars)}

    def vec(self, c):
        return self.p[c]

    def refine(self, matches, strength=2.0):
        by = {}
        for g, c in matches:
            by.setdefault(c, []).append([g["top"], g["bottom"], g["width"]])
        for c, rows in by.items():
            prior, med, n = self.p[c], np.median(rows, axis=0), len(rows)
            for k in range(2):
                prior[k] = (prior[k] * strength + med[k] * n) / (strength + n)
            # a width from a single box could be a merged pair, so it needs a second sample
            if n >= 2:
                prior[2] = med[2]


def _width(protos, c):
    """Learned width in x-heights, or a rough default for characters seen once"""
    w = protos.vec(c)[2]
    if not np.isnan(w):
        return float(w)
    return 0.35 if c in ".,:;'’‘" else 0.9


SIGMA = np.array([0.22, 0.22, 0.30])


def _cost(vec, proto):
    d = (vec - proto) / SIGMA
    d = d[~np.isnan(proto)]
    return float(min((d**2).sum(), 12.0))


def _union(a, b):
    top, bottom = max(a["top"], b["top"]), max(a["bottom"], b["bottom"])
    return np.array([top, bottom, (max(a["x"] + a["w"], b["x"] + b["w"]) - min(a["x"], b["x"])) / a["xh"]])


def _vec(g):
    return np.array([g["top"], g["bottom"], g["width"]])


def align_once(glyphs, chars, protos):
    """Needleman-Wunsch over boxes and characters, with one-box-two-chars and two-boxes-one-char steps"""
    n, m = len(glyphs), len(chars)
    INF = 1e18
    D = np.full((n + 1, m + 1), INF)
    back = {}
    D[0, 0] = 0
    for i in range(n + 1):
        for j in range(m + 1):
            cur = D[i, j]
            if cur >= INF:
                continue

            def go(ni, nj, cost, op):
                if ni <= n and nj <= m and cur + cost < D[ni, nj]:
                    D[ni, nj] = cur + cost
                    back[(ni, nj)] = (i, j, op)

            go(i + 1, j, SKIP_BOX, "skipbox") if i < n else None
            go(i, j + 1, SKIP_CHAR, "skipchar") if j < m else None
            if i < n and j < m:
                go(i + 1, j + 1, _cost(_vec(glyphs[i]), protos.vec(chars[j])), "match")
            if i < n and j + 1 < m:
                a, b = protos.vec(chars[j]), protos.vec(chars[j + 1])
                both = np.array([max(a[0], b[0]), max(a[1], b[1]), np.nan if np.isnan(a[2] + b[2]) else a[2] + b[2]])
                go(i + 1, j + 2, _cost(_vec(glyphs[i]), both) + MERGE_PENALTY, "merge")
            if i + 1 < n and j < m and glyphs[i]["line"] == glyphs[i + 1]["line"]:
                a, b = glyphs[i], glyphs[i + 1]
                gap = (b["x"] - a["x"] - a["w"]) / a["xh"]
                cost = _cost(_union(a, b), protos.vec(chars[j])) + MERGE_PENALTY + min((gap / 0.25) ** 2, 12.0)
                go(i + 2, j + 1, cost, "split")
    path, i, j = [], n, m
    while (i, j) != (0, 0):
        pi, pj, op = back[(i, j)]
        path.append((pi, pj, op))
        i, j = pi, pj
    return path[::-1], D[n, m]


def align(glyphs, text, rounds=3):
    """Labels each glyph dict with char, op and cost. Merged boxes get a string of 2 chars, split boxes share one"""
    chars = [c for c in text if not c.isspace()]
    line_metrics(glyphs)
    protos = Protos(chars)
    for _ in range(rounds):
        path, total = align_once(glyphs, chars, protos)
        matches = [(glyphs[i], chars[j]) for i, j, op in path if op == "match"]
        protos.refine(matches)
    path, total = align_once(glyphs, chars, protos)
    for g in glyphs:
        g["char"], g["align"], g["cost"] = None, "unmatched", None
    for i, j, op in path:
        if op == "match":
            g = glyphs[i]
            g["char"], g["align"] = chars[j], "match"
            g["cost"] = _cost(_vec(g), protos.vec(chars[j]))
        elif op == "merge":
            glyphs[i]["char"], glyphs[i]["align"] = chars[j] + chars[j + 1], "merged"
            wa, wb = (_width(protos, chars[j + k]) for k in (0, 1))
            glyphs[i]["split_at"] = float(wa / (wa + wb))
        elif op == "split":
            for k in (i, i + 1):
                glyphs[k]["char"], glyphs[k]["align"] = chars[j], "split"
    return total


def card_targets(xlsx, name, font):
    """Returns the glyph characters the Trace targets tab asks for from a card in one font"""
    ws = openpyxl.load_workbook(xlsx)["Trace targets"]
    return [r[4] for r in ws.iter_rows(min_row=2, values_only=True) if r[2] == name and r[3] == font]
